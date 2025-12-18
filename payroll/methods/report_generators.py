"""
report_generators.py

Report generation functions for various report types
"""

import csv
from datetime import datetime
from io import BytesIO
from decimal import Decimal

from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from reportlab.pdfgen import canvas
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

from payroll.models.models import Payslip, Contract
from employee.models import Employee
from base.models import Company, Department


def generate_payroll_register_pdf(parameters, generated_report):
    """
    Generate Payroll Register in PDF format
    """
    if not PDF_AVAILABLE:
        raise ImportError("ReportLab is required for PDF generation. Install with: pip install reportlab")
    
    date_from = parameters.get('date_from')
    date_to = parameters.get('date_to')
    company = parameters.get('company')
    department = parameters.get('department')
    
    # Get payslips
    payslips = Payslip.objects.filter(
        start_date__gte=date_from,
        end_date__lte=date_to,
        status='confirmed'
    )
    
    if company:
        payslips = payslips.filter(employee__employee_work_info__company=company)
    if department:
        payslips = payslips.filter(employee__employee_work_info__department=department)
    
    # Create PDF
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )
    
    elements = []
    styles = getSampleStyleSheet()
    
    # Title
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor('#1a472a'),
        spaceAfter=30,
        alignment=1  # Center
    )
    
    title = Paragraph(f"Payroll Register<br/>{date_from} to {date_to}", title_style)
    elements.append(title)
    elements.append(Spacer(1, 12))
    
    # Prepare table data
    table_data = [
        ['Employee ID', 'Employee Name', 'Department', 'Basic Pay', 'Allowances', 
         'Gross Pay', 'Deductions', 'Net Pay']
    ]
    
    total_basic = Decimal('0')
    total_gross = Decimal('0')
    total_deductions = Decimal('0')
    total_net = Decimal('0')
    
    for payslip in payslips:
        employee_work_info = payslip.employee.employee_work_info.first()
        dept_name = employee_work_info.department.department if employee_work_info and employee_work_info.department else 'N/A'
        
        allowances = payslip.gross_pay - payslip.basic_pay
        
        table_data.append([
            payslip.employee.employee_id,
            str(payslip.employee),
            dept_name,
            f"₱{payslip.basic_pay:,.2f}",
            f"₱{allowances:,.2f}",
            f"₱{payslip.gross_pay:,.2f}",
            f"₱{payslip.deduction:,.2f}",
            f"₱{payslip.net_pay:,.2f}",
        ])
        
        total_basic += payslip.basic_pay
        total_gross += payslip.gross_pay
        total_deductions += payslip.deduction
        total_net += payslip.net_pay
    
    # Add totals row
    table_data.append([
        '', 'TOTAL', '',
        f"₱{total_basic:,.2f}",
        '',
        f"₱{total_gross:,.2f}",
        f"₱{total_deductions:,.2f}",
        f"₱{total_net:,.2f}",
    ])
    
    # Create table
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a472a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#f0f0f0')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('ALIGN', (3, 1), (-1, -1), 'RIGHT'),
    ]))
    
    elements.append(table)
    
    # Build PDF
    doc.build(elements)
    
    # Prepare response
    buffer.seek(0)
    filename = f"payroll_register_{date_from}_to_{date_to}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    return response


def generate_payroll_register_excel(parameters, generated_report):
    """
    Generate Payroll Register in Excel format
    """
    if not EXCEL_AVAILABLE:
        raise ImportError("openpyxl is required for Excel generation. Install with: pip install openpyxl")
    
    date_from = parameters.get('date_from')
    date_to = parameters.get('date_to')
    company = parameters.get('company')
    department = parameters.get('department')
    
    # Get payslips
    payslips = Payslip.objects.filter(
        start_date__gte=date_from,
        end_date__lte=date_to,
        status='confirmed'
    )
    
    if company:
        payslips = payslips.filter(employee__employee_work_info__company=company)
    if department:
        payslips = payslips.filter(employee__employee_work_info__department=department)
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll Register"
    
    # Styling
    header_fill = PatternFill(start_color="1a472a", end_color="1a472a", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Title
    ws.merge_cells('A1:H1')
    ws['A1'] = f"Payroll Register - {date_from} to {date_to}"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')
    
    # Headers
    headers = ['Employee ID', 'Employee Name', 'Department', 'Basic Pay', 
               'Allowances', 'Gross Pay', 'Deductions', 'Net Pay']
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Data rows
    row_num = 4
    total_basic = Decimal('0')
    total_gross = Decimal('0')
    total_deductions = Decimal('0')
    total_net = Decimal('0')
    
    for payslip in payslips:
        employee_work_info = payslip.employee.employee_work_info.first()
        dept_name = employee_work_info.department.department if employee_work_info and employee_work_info.department else 'N/A'
        
        allowances = payslip.gross_pay - payslip.basic_pay
        
        ws.cell(row=row_num, column=1).value = payslip.employee.employee_id
        ws.cell(row=row_num, column=2).value = str(payslip.employee)
        ws.cell(row=row_num, column=3).value = dept_name
        ws.cell(row=row_num, column=4).value = float(payslip.basic_pay)
        ws.cell(row=row_num, column=5).value = float(allowances)
        ws.cell(row=row_num, column=6).value = float(payslip.gross_pay)
        ws.cell(row=row_num, column=7).value = float(payslip.deduction)
        ws.cell(row=row_num, column=8).value = float(payslip.net_pay)
        
        # Apply borders
        for col in range(1, 9):
            ws.cell(row=row_num, column=col).border = border
            if col >= 4:  # Number columns
                ws.cell(row=row_num, column=col).number_format = '#,##0.00'
        
        total_basic += payslip.basic_pay
        total_gross += payslip.gross_pay
        total_deductions += payslip.deduction
        total_net += payslip.net_pay
        
        row_num += 1
    
    # Totals row
    ws.cell(row=row_num, column=2).value = "TOTAL"
    ws.cell(row=row_num, column=2).font = Font(bold=True)
    ws.cell(row=row_num, column=4).value = float(total_basic)
    ws.cell(row=row_num, column=6).value = float(total_gross)
    ws.cell(row=row_num, column=7).value = float(total_deductions)
    ws.cell(row=row_num, column=8).value = float(total_net)
    
    for col in range(1, 9):
        ws.cell(row=row_num, column=col).border = border
        ws.cell(row=row_num, column=col).font = Font(bold=True)
        if col >= 4:
            ws.cell(row=row_num, column=col).number_format = '#,##0.00'
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['H'].width = 15
    
    # Save to buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"payroll_register_{date_from}_to_{date_to}.xlsx"
    generated_report.file_name = filename
    
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    return response


def generate_variance_report(parameters, generated_report):
    """
    Generate variance report comparing current vs previous period
    """
    # Placeholder implementation
    date_from = parameters.get('date_from')
    date_to = parameters.get('date_to')
    
    filename = f"variance_report_{date_from}_to_{date_to}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"Variance Report - Implementation Pending")
    
    return response


def generate_employer_contributions(parameters, generated_report):
    """
    Generate employer contributions report
    """
    # Placeholder - to be implemented with actual contribution calculations
    filename = f"employer_contributions_{parameters.get('date_from')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"Employer Contributions Report - Implementation Pending")
    
    return response


def generate_bank_advice_list(parameters, generated_report):
    """
    Generate bank advice list
    """
    filename = f"bank_advice_list_{parameters.get('date_from')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"Bank Advice List - Implementation Pending")
    
    return response


def generate_bank_file(parameters, generated_report):
    """
    Generate bank file for salary transfer
    """
    date_from = parameters.get('date_from')
    date_to = parameters.get('date_to')
    
    # Get payslips
    payslips = Payslip.objects.filter(
        start_date__gte=date_from,
        end_date__lte=date_to,
        status='confirmed'
    )
    
    # Create CSV
    response = HttpResponse(content_type='text/csv')
    filename = f"bank_file_{date_from}_to_{date_to}.csv"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    writer = csv.writer(response)
    writer.writerow(['Employee ID', 'Account Number', 'Employee Name', 'Net Pay'])
    
    for payslip in payslips:
        employee = payslip.employee
        bank_account = getattr(employee, 'bank_account_no', 'N/A')
        writer.writerow([
            employee.employee_id,
            bank_account,
            str(employee),
            f"{payslip.net_pay:.2f}"
        ])
    
    generated_report.file_name = filename
    return response


# Statutory reports
def generate_sss_report(parameters, generated_report):
    """Generate SSS contribution report"""
    filename = f"sss_report_{parameters.get('date_from')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"SSS Report - Implementation Pending")
    return response


def generate_philhealth_report(parameters, generated_report):
    """Generate PhilHealth contribution report"""
    filename = f"philhealth_report_{parameters.get('date_from')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"PhilHealth Report - Implementation Pending")
    return response


def generate_pagibig_report(parameters, generated_report):
    """Generate PAG-IBIG contribution report"""
    filename = f"pagibig_report_{parameters.get('date_from')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"PAG-IBIG Report - Implementation Pending")
    return response


def generate_bir_report(parameters, generated_report):
    """Generate BIR withholding tax report"""
    filename = f"bir_report_{parameters.get('date_from')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"BIR Report - Implementation Pending")
    return response


def generate_net_pay_validation(parameters, generated_report):
    """Generate net pay validation report"""
    filename = f"net_pay_validation_{parameters.get('date_from')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"Net Pay Validation - Implementation Pending")
    return response


def generate_basic_salary_report(parameters, generated_report):
    """Generate basic salary report"""
    filename = f"basic_salary_report_{parameters.get('date_from')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"Basic Salary Report - Implementation Pending")
    return response


def generate_withholding_tax_validation(parameters, generated_report):
    """Generate withholding tax validation report"""
    filename = f"withholding_tax_validation_{parameters.get('date_from')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"Withholding Tax Validation - Implementation Pending")
    return response


def generate_certificate_contribution(parameters, generated_report):
    """Generate certificate of contribution"""
    filename = f"certificate_contribution_{parameters.get('employee')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"Certificate of Contribution - Implementation Pending")
    return response


def generate_demographic_report(parameters, generated_report):
    """Generate demographic report"""
    filename = f"demographic_report_{parameters.get('date_from')}.pdf"
    generated_report.file_name = filename
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(b"Demographic Report - Implementation Pending")
    return response
